import pytest
from agentic_rag.app.document_parser import DocumentParser
from agentic_rag.app.chunking_strategy import chunk_text

# Unit tests for the Document Parser

def create_txt_file(tmp_path, content="Hello World"):
    file_path = tmp_path / "sample.txt"
    with open(file_path, "w", encoding="utf-8") as f:
        f.write(content)
    return file_path

def test_parse_txt_file(tmp_path):
    parser = DocumentParser()
    file_path = create_txt_file(tmp_path, "Test text file content.")
    # Consume the generator to get a list of results
    results = list(parser.parse(str(file_path)))
    assert len(results) == 1
    result = results[0]
    assert "text" in result
    assert result["text"] == "Test text file content."
    assert "metadata" in result
    assert result["metadata"]["file_name"] == "sample.txt"
    assert result["metadata"]["file_size"] > 0

def test_parse_nonexistent_file(tmp_path):
    parser = DocumentParser()
    fake_path = tmp_path / "nofile.txt"
    with pytest.raises(FileNotFoundError):
        # Consume the generator inside the 'raises' block to trigger the error
        list(parser.parse(str(fake_path)))

def test_parse_unsupported_file_type(tmp_path):
    parser = DocumentParser()
    file_path = tmp_path / "sample.xyz"
    file_path.write_text("dummy")
    with pytest.raises(ValueError):
        # Consume the generator to trigger the error
        list(parser.parse(str(file_path)))

def test_parse_pdf_file(tmp_path):
    parser = DocumentParser()
    try:
        from fpdf import FPDF
    except ImportError:
        pytest.skip("fpdf not installed")
    pdf_path = tmp_path / "sample.pdf"
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, txt="PDF Content", ln=True)
    pdf.output(str(pdf_path))
    results = list(parser.parse(str(pdf_path)))
    assert len(results) == 1
    result = results[0]
    assert "PDF Content" in result["text"]
    assert result["metadata"]["file_name"] == "sample.pdf"

def test_parse_docx_file(tmp_path):
    parser = DocumentParser()
    try:
        import docx
    except ImportError:
        pytest.skip("python-docx not installed")
    docx_path = tmp_path / "sample.docx"
    doc = docx.Document()
    doc.add_paragraph("DOCX Content")
    doc.save(str(docx_path))
    results = list(parser.parse(str(docx_path)))
    assert len(results) == 1
    result = results[0]
    assert "DOCX Content" in result["text"]
    assert result["metadata"]["file_name"] == "sample.docx"

def test_document_parser_txt(tmp_path):
    """Tests parsing a simple .txt file."""
    parser = DocumentParser()
    d = tmp_path / "sub"
    d.mkdir()
    p = d / "hello.txt"
    p.write_text("This is a test.")
    
    results = list(parser.parse(str(p)))
    assert len(results) == 1
    result = results[0]
    
    assert "text" in result
    assert "metadata" in result
    assert result["text"] == "This is a test."
    assert result["metadata"]["file_name"] == "hello.txt"

# Unit tests for the Chunking Strategy (These tests are fine and need no changes)
def test_chunk_text_returns_list_of_strings():
    """Tests that the chunker returns a list of strings."""
    sample_text = "This is a long text that needs to be chunked." * 100
    chunks = chunk_text(sample_text)
    
    assert isinstance(chunks, list)
    assert len(chunks) > 1
    assert isinstance(chunks[0], str)

def test_chunk_text_handles_short_text():
    """Tests that the chunker handles text shorter than the chunk size."""
    sample_text = "This is a short text."
    chunks = chunk_text(sample_text)
    
    assert len(chunks) == 1
    assert chunks[0] == sample_text

def test_chunk_text_basic():
    text = "a" * 1200
    chunks = chunk_text(text, chunk_size=500, chunk_overlap=50)
    assert isinstance(chunks, list)
    assert all(isinstance(chunk, str) for chunk in chunks)
    assert len(chunks) == 3
    assert chunks[0] == "a" * 500
    assert chunks[1] == "a" * 500
    assert chunks[2] == "a" * 300

def test_chunk_text_overlap():
    text = "abcdefghijklmnopqrstuvwxyz" * 2  # 52 chars
    chunks = chunk_text(text, chunk_size=20, chunk_overlap=5)
    assert len(chunks) == 4
    assert chunks[0][15:] == chunks[1][:5]  # overlap
    assert chunks[1][15:] == chunks[2][:5]  # overlap

def test_chunk_text_small_text():
    text = "short text"
    chunks = chunk_text(text, chunk_size=100, chunk_overlap=10)
    assert len(chunks) == 1
    assert chunks[0] == text

def test_chunk_text_invalid_input():
    with pytest.raises(ValueError):
        chunk_text(12345)